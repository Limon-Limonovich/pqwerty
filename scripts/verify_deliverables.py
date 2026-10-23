from pathlib import Path
from datetime import datetime
import math, zipfile
import openpyxl
from docx import Document
P=Path(__file__).resolve().parents[1]/'Результат'
f=openpyxl.load_workbook(P/'Практические_задания.xlsx',data_only=False)
v=openpyxl.load_workbook(P/'Практические_задания.xlsx',data_only=True)
assert len(f.sheetnames)==18
assert sum(len(s._charts) for s in f)==13
count=0
for s in f:
 for row in s:
  for c in row:
   if c.data_type=='f':
    count+=1
    cache=v[s.title][c.coordinate]
    assert cache.value is not None or (s.title=='07_Дедлайн' and c.coordinate=='E1'),(s.title,c.coordinate)
    assert cache.data_type!='e',(s.title,c.coordinate,cache.value)
for s in list(f)[13:]:
 assert len(s._charts)==2,s.title
 assert len({type(c).__name__ for c in s._charts})==2,s.title

def check(sn,addr,expected):
 got=v[sn][addr].value
 assert math.isclose(got,expected,rel_tol=1e-10) if isinstance(expected,(int,float)) else got==expected,(sn,addr,got,expected)
for r in range(3,7):
 s=v['01_ТекстильПром'];check(s.title,f'F{r}',s[f'E{r}'].value*100/s[f'D{r}'].value);check(s.title,f'J{r}',(s[f'H{r}'].value-s[f'G{r}'].value)*s[f'E{r}'].value)
check('01_ТекстильПром','J7',1094000)
check('02_Сумма_среднее','D3',550);check('02_Сумма_среднее','D4',55)
check('03_Среднее_без_нулей','B1',40)
check('04_Скидка','B7',4750);check('04_Скидка','B8',4500.9)
check('05_Доступ','C7','Отказано');check('05_Доступ','C8','Доступ разрешен');check('05_Доступ','C9','Отказано')
check('06_Срок_проекта','C1',datetime(2024,10,15));check('07_Дедлайн','C10',datetime(2024,10,11))
assert len(f['07_Дедлайн'].conditional_formatting)==2
check('08_Продажи','E2',176.25);check('08_Продажи','E4','Декабрь');check('08_Продажи','E6',6)
check('09_Себестоимость','B9',1260000)
for r,e in enumerate(['Годна','Брак','Годна','Годна','Брак'],2):check('10_Качество',f'D{r}',e)
check('11_Заказ_праздники','D2',datetime(2025,7,3))
check('12_Партии_ткани','B10','Партия 1');check('12_Партии_ткани','B12',math.sqrt(250))
check('13_В1_Рабочее_время','F4',7.5);check('13_В1_Рабочее_время','B9',14880)
check('14_В2_Склад','B9',90);check('15_В3_Производство','B7',1652)
check('16_В4_Доставка','B9',15686.5);check('17_В5_Показатели_цеха','B10',1061900)
for sn in ['13_В1_Рабочее_время','14_В2_Склад','15_В3_Производство','16_В4_Доставка','17_В5_Показатели_цеха']:
 assert len(f[sn]._charts)==2
for name in ['Практические_задания.xlsx','Отчёт_по_уроку.docx']:
 with zipfile.ZipFile(P/name) as z:assert z.testzip() is None
report=Document(P/'Отчёт_по_уроку.docx')
assert len(report.inline_shapes)==2
assert len(report.tables)==4
print(f'OK: 18 листов, 13 диаграмм, {count} формул с сохранёнными результатами; арифметика, даты, граничные случаи и структура DOCX проверены.')
print('Проверка программная. Пересчёт в настольном Excel и визуальная проверка Word не выполнялись.')
